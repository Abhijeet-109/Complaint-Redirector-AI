"""Focused tests for durable chatbot audit records and Excel export."""

import os
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import uuid
from unittest.mock import patch

os.environ.setdefault("DATABASE_URL", "sqlite://")

from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.Database.connection import Base
from backend.Database.models import ChatbotAuditRequest, ChatbotAuditStage
from backend.chatbot import audit_logging


class ChatbotAuditLoggingTests(unittest.TestCase):
    def setUp(self):
        self.test_root = Path(__file__).resolve().parents[2]
        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        self.factory = sessionmaker(bind=self.engine)
        self.workbook_path = self.test_root / f".audit-test-{uuid.uuid4().hex}.xlsx"
        self.session_patch = patch.object(audit_logging, "SessionLocal", self.factory)
        self.path_patch = patch.object(audit_logging, "WORKBOOK_PATH", self.workbook_path)
        self.session_patch.start()
        self.path_patch.start()

    def tearDown(self):
        self.path_patch.stop()
        self.session_patch.stop()
        Base.metadata.drop_all(self.engine)
        self.engine.dispose()
        for path in self.test_root.glob(f"{self.workbook_path.stem}*"):
            path.unlink(missing_ok=True)

    def _persist(self, query="Email me at person@example.com or +1 202 555 0142"):
        trace = audit_logging.Trace(query, "session-123", user_id=7)
        with trace.stage("input_guardrail") as stage:
            stage.decision = "allow"
        with trace.stage("retrieval") as stage:
            stage.outcome("fallback", decision="empty_result", category="no_document_match")
        audit_logging.persist_audit(
            trace,
            answer="Contact person@example.com",
            status="fallback",
            latency_ms=12.5,
            fallback=True,
        )
        return trace

    def test_sensitive_data_is_redacted_before_database_persistence(self):
        self._persist()
        db = self.factory()
        try:
            row = db.query(ChatbotAuditRequest).one()
            self.assertNotIn("person@example.com", row.query_text)
            self.assertNotIn("202 555 0142", row.query_text)
            self.assertNotIn("person@example.com", row.answer_text)
            self.assertEqual(row.status, "fallback")
            self.assertEqual(db.query(ChatbotAuditStage).count(), 3)
        finally:
            db.close()

    def test_excel_export_appends_headers_and_is_idempotent(self):
        self._persist("=HYPERLINK(\"https://bad.example\")")
        self.assertEqual(audit_logging.export_pending_batch(), 1)
        self.assertEqual(audit_logging.export_pending_batch(), 0)
        wb = load_workbook(self.workbook_path, read_only=True, data_only=True)
        try:
            requests = wb["Request Logs"]
            rows = list(requests.iter_rows(values_only=True))
            self.assertEqual(rows[0][0], "request_id")
            self.assertEqual(len(rows), 2)
            self.assertTrue(rows[1][4].startswith("'="))
            self.assertEqual(len(list(wb["Stage Logs"].iter_rows(values_only=True))), 4)
            self.assertGreater(wb["Summary"].max_row, 4)
        finally:
            wb.close()

    def test_export_creates_missing_workbook_directory_before_locking(self):
        self._persist("A normal question")
        temporary_root = self.test_root / f".audit-test-dir-{uuid.uuid4().hex}"
        workbook_path = temporary_root / "nested" / "logs" / "audit.xlsx"
        try:
            with patch.object(audit_logging, "WORKBOOK_PATH", workbook_path):
                self.assertEqual(audit_logging.export_pending_batch(), 1)
                self.assertTrue(workbook_path.is_file())
        finally:
            workbook_path.unlink(missing_ok=True)
            workbook_path.with_suffix(workbook_path.suffix + ".lock").unlink(missing_ok=True)
            workbook_path.parent.rmdir()
            workbook_path.parent.parent.rmdir()
            temporary_root.rmdir()

    def test_concurrent_export_calls_do_not_duplicate_request_rows(self):
        self._persist("A normal question")
        with ThreadPoolExecutor(max_workers=2) as pool:
            list(pool.map(lambda _n: audit_logging.export_pending_batch(), range(2)))
        wb = load_workbook(self.workbook_path, read_only=True)
        try:
            self.assertEqual(wb["Request Logs"].max_row, 2)
        finally:
            wb.close()

    def test_malformed_workbook_is_preserved_and_reinitialized(self):
        self.workbook_path.write_bytes(b"not an xlsx file")
        self._persist("A normal question")
        self.assertEqual(audit_logging.export_pending_batch(), 1)
        self.assertTrue(list(self.test_root.glob(f"{self.workbook_path.stem}.damaged-*.xlsx")))
        wb = load_workbook(self.workbook_path, read_only=True)
        try:
            self.assertTrue(wb.sheetnames)
        finally:
            wb.close()

    def test_excel_failure_keeps_postgres_record_pending_and_retry_recovers(self):
        trace = self._persist("A normal question")
        with patch.object(audit_logging, "_workbook", side_effect=OSError("disk unavailable")):
            self.assertEqual(audit_logging.export_pending_batch(), 0)
        db = self.factory()
        try:
            row = db.query(ChatbotAuditRequest).filter_by(request_id=trace.request_id).one()
            export_stage = db.query(ChatbotAuditStage).filter_by(
                request_id=trace.request_id, stage_name="excel_export",
            ).one()
            self.assertIsNone(row.exported_at)
            self.assertEqual(row.export_error, "OSError")
            self.assertEqual(export_stage.status, "failed")
        finally:
            db.close()
        self.assertEqual(audit_logging.export_pending_batch(), 1)


if __name__ == "__main__":
    unittest.main()
