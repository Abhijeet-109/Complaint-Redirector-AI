
import os

import pandas as pd
import pytesseract

from docx import Document as WordDocument
from pypdf import PdfReader
from pdf2image import convert_from_path
from PIL import Image


def load_document(file_path):

    extension = os.path.splitext(file_path)[1].lower()
    filename = os.path.basename(file_path)

    documents = []

    # PDF: extract text and use OCR for scanned pages
    if extension == ".pdf":

        reader = PdfReader(file_path)
        pages = reader.pages

        # Extract text from every PDF page
        page_texts = [
            page.extract_text() or ""
            for page in pages
        ]

        # Convert pages to images only if OCR is needed
        if any(not text.strip() for text in page_texts):

            images = convert_from_path(file_path)

            for i, text in enumerate(page_texts):

                if not text.strip():
                    page_texts[i] = pytesseract.image_to_string(
                        images[i]
                    )

        for i, text in enumerate(page_texts):
            if text.strip():
                documents.append({
                    "page_content": text,
                    "metadata": {
                        "source": filename,
                        "page": i + 1
                    }
                })

    # Word document
    elif extension == ".docx":

        word_doc = WordDocument(file_path)

        text_parts = [
            paragraph.text
            for paragraph in word_doc.paragraphs
            if paragraph.text.strip()
        ]

        # Include text from Word tables
        for table in word_doc.tables:
            for row in table.rows:
                text_parts.append(
                    " | ".join(cell.text for cell in row.cells)
                )

        text = "\n".join(text_parts)

        documents.append({
            "page_content": text,
            "metadata": {"source": filename}
        })

    # Text file
    elif extension == ".txt":

        with open(file_path, "r", encoding="utf-8-sig") as file:
            text = file.read()

        documents.append({
            "page_content": text,
            "metadata": {"source": filename}
        })

    # Excel and CSV files
    elif extension in [".xlsx", ".xls", ".csv"]:

        if extension == ".csv":
            sheets = {"CSV": pd.read_csv(file_path)}
        else:
            sheets = pd.read_excel(
                file_path,
                sheet_name=None
            )

        for sheet_name, dataframe in sheets.items():

            text = dataframe.to_string(index=False)

            documents.append({
                "page_content": f"Sheet: {sheet_name}\n{text}",
                "metadata": {
                    "source": filename,
                    "sheet": sheet_name
                }
            })

    # Image files
    elif extension in [".png", ".jpg", ".jpeg", ".tiff"]:

        image = Image.open(file_path)

        text = pytesseract.image_to_string(image)

        documents.append({
            "page_content": text,
            "metadata": {"source": filename}
        })

    else:
        raise ValueError(
            f"Unsupported file format: {extension}"
        )

    # Remove empty results
    documents = [
        doc for doc in documents
        if doc["page_content"].strip()
    ]

    if not documents:
        raise ValueError(
            f"No readable text found in {filename}"
        )

    return documents
