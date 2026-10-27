"""Chatbot request tracing, privacy filtering, and durable Excel export."""

import asyncio
import contextvars
import logging
import os
import re
import time
import threading
import uuid
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func

from backend.Database.connection import SessionLocal
from backend.Database.models import ChatbotAuditRequest, ChatbotAuditStage

logger = logging.getLogger("flatkart.chatbot.audit")
DEFAULT_WORKBOOK = Path(__file__).resolve().parents[1] / "logs" / "chatbot_audit.xlsx"
WORKBOOK_PATH = Path(os.getenv("CHATBOT_AUDIT_WORKBOOK", str(DEFAULT_WORKBOOK))).expanduser()
EXPORT_BATCH_SIZE = max(1, min(int(os.getenv("CHATBOT_AUDIT_EXPORT_BATCH_SIZE", "100")), 1000))
EXPORT_POLL_SECONDS = max(1, int(os.getenv("CHATBOT_AUDIT_EXPORT_POLL_SECONDS", "5")))

REQUEST_HEADERS = [
    "request_id", "timestamp_utc", "session_id", "user_id", "query", "answer", "status",
    "failure_stage", "failure_category", "failure_reason", "http_status", "latency_ms",
    "source_count", "source_names", "model_name", "fallback", "fallback_reason", "exported_at_utc",
    "export_attempts", "export_failures", "export_error",
]
STAGE_HEADERS = [
    "request_id", "stage_name", "status", "started_at_utc", "completed_at_utc", "duration_ms",
    "decision", "failure_category", "failure_reason", "exception_type", "retry_count", "skip_reason",
]

_EMAIL = re.compile(r"\b[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}\b")
_PHONE = re.compile(r"(?<!\w)(?:\+?\d[\d ().-]{7,}\d)(?!\w)")
_SECRET_ASSIGNMENT = re.compile(
    r"(?i)\b(?:api[_ -]?key|access[_ -]?token|refresh[_ -]?token|oauth[_ -]?token|"
    r"password|passwd|client[_ -]?secret|authorization|secret)\s*[:=]\s*[^\s,;]+"
)
_TOKEN = re.compile(
    r"(?i)\b(?:bearer\s+\S+|(?:sk-[A-Za-z0-9_-]{12,}|gh[pousr]_[A-Za-z0-9_]{20,})|"
    r"eyJ[A-Za-z0-9_-]{10,}\.[A-Za-z0-9_-]{10,}\.[A-Za-z0-9_-]{10,})\b"
)
_CURRENT_TRACE = contextvars.ContextVar("chatbot_audit_trace", default=None)
_PROCESS_EXPORT_LOCK = threading.Lock()


def utcnow():
    return datetime.now(timezone.utc)


def _utc_iso(value):
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=timezone.utc)
    return value.astimezone(timezone.utc).isoformat()


def redact(value, limit=None):
    """Redact common direct identifiers and credential-like strings."""
    if value is None:
        return None
    text = _SECRET_ASSIGNMENT.sub("[redacted]", str(value))
    text = _TOKEN.sub("[redacted]", text)
    text = _EMAIL.sub("[email]", text)
    text = _PHONE.sub("[phone]", text)
    if limit is not None:
        text = text[:limit]
    return text


def _safe_cell(value):
    """Force user-controlled strings to literal text in Excel."""
    if not isinstance(value, str):
        return value
    safe = redact(value)
    if safe.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + safe
    return safe


class Stage:
    def __init__(self, trace, name, decision=None):
        self.trace, self.name = trace, name
        trace.active_stage = self
        self.started_at = utcnow()
        self.started = time.perf_counter()
        self.status, self.decision = "success", decision
        self.failure_category = self.failure_reason = self.exception_type = self.skip_reason = None

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, _tb):
        if exc_type is not None:
            self.outcome("failed", category="infrastructure_error", reason="Stage failed.", exception=exc_type.__name__)
        self.finish()
        return False

    def outcome(self, status, decision=None, category=None, reason=None, exception=None):
        self.status = status
        self.decision = decision
        self.failure_category = category
        self.failure_reason = redact(reason, 240) if reason else None
        self.exception_type = exception

    def skip(self, reason):
        self.outcome("skipped", decision="skipped", reason=reason)
        self.skip_reason = redact(reason, 240)

    def finish(self):
        self.trace.stages.append({
            "stage_name": self.name,
            "status": self.status,
            "started_at": self.started_at,
            "completed_at": utcnow(),
            "duration_ms": (time.perf_counter() - self.started) * 1000,
            "decision": redact(self.decision, 120) if self.decision else None,
            "failure_category": self.failure_category,
            "failure_reason": self.failure_reason,
            "exception_type": self.exception_type,
            "retry_count": 0,
            "skip_reason": self.skip_reason,
        })
        self.trace.active_stage = None


class Trace:
    def __init__(self, question, session_id, user_id=None):
        self.request_id = str(uuid.uuid4())
        self.created_at = utcnow()
        self.question = redact(question)
        self.session_id = redact(session_id, 100) or ""
        self.user_id = str(user_id)[:64] if user_id is not None else None
        self.stages = []
        self.active_stage = None

    def stage(self, name, decision=None):
        return Stage(self, name, decision)


def current_trace():
    return _CURRENT_TRACE.get()


def activate(trace):
    return _CURRENT_TRACE.set(trace)


def deactivate(token):
    _CURRENT_TRACE.reset(token)


def persist_audit(trace, *, answer, status, failure_stage=None, failure_category=None,
                  failure_reason=None, http_status=200, latency_ms=0, sources=None,
                  fallback=False, fallback_reason=None):
    """Persist request and stages in one PostgreSQL transaction before response."""
    db = SessionLocal()
    try:
        request = ChatbotAuditRequest(
            request_id=trace.request_id,
            created_at=trace.created_at,
            session_id=trace.session_id,
            user_id=trace.user_id,
            query_text=trace.question,
            answer_text=redact(answer) if answer is not None else None,
            status=status,
            failure_stage=failure_stage,
            failure_category=failure_category,
            failure_reason=redact(failure_reason, 240) if failure_reason else None,
            http_status=http_status,
            latency_ms=latency_ms,
            source_count=len(sources or []),
            source_names=list(dict.fromkeys(redact(str(item.get("document", "")), 180) for item in (sources or []) if item.get("document"))),
            model_name="openai/gpt-oss-120b",
            fallback=bool(fallback),
            fallback_reason=redact(fallback_reason, 240) if fallback_reason else None,
        )
        db.add(request)
        db.add_all([ChatbotAuditStage(request_id=trace.request_id, **row) for row in trace.stages])
        export_time = utcnow()
        db.add(ChatbotAuditStage(
            request_id=trace.request_id,
            stage_name="excel_export",
            status="pending",
            started_at=export_time,
            completed_at=export_time,
            duration_ms=0,
            decision="queued",
            retry_count=0,
        ))
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()


def _init_sheet(ws, headers):
    if ws.cell(1, 1).value is None:
        for column, header in enumerate(headers, 1):
            ws.cell(row=1, column=column, value=header)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243B53")


def _format_sheet(ws):
    if ws.title == "Request Logs":
        widths = [38, 28, 24, 12, 48, 48, 18, 22, 24, 42, 14, 16, 14, 32, 28, 12, 42, 28, 16, 16, 20]
    elif ws.title == "Stage Logs":
        widths = [38, 26, 18, 28, 28, 16, 24, 26, 42, 24, 14, 42]
    else:
        widths = [40, 32]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _workbook():
    WORKBOOK_PATH.parent.mkdir(parents=True, exist_ok=True)
    if WORKBOOK_PATH.exists():
        try:
            wb = load_workbook(WORKBOOK_PATH)
        except Exception:
            # Preserve malformed workbook as evidence; write to a recoverable new file.
            damaged = WORKBOOK_PATH.with_name(f"{WORKBOOK_PATH.stem}.damaged-{int(time.time())}{WORKBOOK_PATH.suffix}")
            WORKBOOK_PATH.replace(damaged)
            wb = Workbook()
    else:
        wb = Workbook()
    for title in ("Request Logs", "Stage Logs", "Summary"):
        if title not in wb.sheetnames:
            wb.create_sheet(title)
    default = wb["Sheet"] if "Sheet" in wb.sheetnames else None
    if default and default.max_row == 1 and default.cell(1, 1).value is None:
        wb.remove(default)
    _init_sheet(wb["Request Logs"], REQUEST_HEADERS)
    _init_sheet(wb["Stage Logs"], STAGE_HEADERS)
    return wb


def _lock_file():
    # The lock file is a sibling of the workbook, so create its parent before
    # opening it. `_workbook()` also creates this directory, but the lock is
    # acquired first during export.
    WORKBOOK_PATH.parent.mkdir(parents=True, exist_ok=True)
    lock_path = WORKBOOK_PATH.with_suffix(WORKBOOK_PATH.suffix + ".lock")
    handle = open(lock_path, "a+b")
    if os.name == "nt":
        import msvcrt
        if os.fstat(handle.fileno()).st_size == 0:
            handle.write(b"0")
            handle.flush()
        handle.seek(0)
        msvcrt.locking(handle.fileno(), msvcrt.LK_LOCK, 1)
    else:
        import fcntl
        fcntl.flock(handle.fileno(), fcntl.LOCK_EX)
    return handle


def _unlock_file(handle):
    try:
        if os.name == "nt":
            import msvcrt
            handle.seek(0)
            msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
        else:
            import fcntl
            fcntl.flock(handle.fileno(), fcntl.LOCK_UN)
    finally:
        handle.close()


def _refresh_summary(wb, db):
    ws = wb["Summary"]
    ws.delete_rows(1, ws.max_row)
    total = db.query(func.count(ChatbotAuditRequest.id)).scalar() or 0
    latency_rows = db.query(ChatbotAuditRequest.latency_ms).order_by(ChatbotAuditRequest.created_at.desc()).limit(10000).all()
    latencies = sorted(value[0] for value in latency_rows)
    def percentile(fraction):
        if not latencies:
            return None
        return round(latencies[min(int((len(latencies) - 1) * fraction), len(latencies) - 1)], 1)
    rows = [
        ("Metric", "Value"),
        ("Total requests", total),
        ("Successful requests", db.query(func.count(ChatbotAuditRequest.id)).filter(ChatbotAuditRequest.status == "success").scalar() or 0),
        ("Blocked requests", db.query(func.count(ChatbotAuditRequest.id)).filter(ChatbotAuditRequest.status == "blocked").scalar() or 0),
        ("Fallback responses", db.query(func.count(ChatbotAuditRequest.id)).filter(ChatbotAuditRequest.fallback.is_(True)).scalar() or 0),
        ("Failed requests", db.query(func.count(ChatbotAuditRequest.id)).filter(ChatbotAuditRequest.status == "failed").scalar() or 0),
        ("Average latency ms", round(db.query(func.avg(ChatbotAuditRequest.latency_ms)).scalar() or 0, 1)),
        ("Latency p50 ms (up to 10000 rows)", percentile(0.50)),
        ("Latency p90 ms (up to 10000 rows)", percentile(0.90)),
        ("Latency p95 ms (up to 10000 rows)", percentile(0.95)),
        ("Excel export failures", db.query(func.sum(ChatbotAuditRequest.export_failures)).scalar() or 0),
        ("Input guardrail rejections", db.query(func.count(ChatbotAuditStage.id)).filter(ChatbotAuditStage.stage_name == "input_guardrail", ChatbotAuditStage.status == "blocked").scalar() or 0),
        ("Output guardrail rejections", db.query(func.count(ChatbotAuditStage.id)).filter(ChatbotAuditStage.stage_name == "output_guardrail", ChatbotAuditStage.status == "blocked").scalar() or 0),
        ("Retrieval failures", db.query(func.count(ChatbotAuditStage.id)).filter(ChatbotAuditStage.stage_name == "retrieval", ChatbotAuditStage.status == "failed").scalar() or 0),
        ("Empty retrieval fallbacks", db.query(func.count(ChatbotAuditStage.id)).filter(ChatbotAuditStage.stage_name == "retrieval", ChatbotAuditStage.status == "fallback").scalar() or 0),
        ("Last successful export UTC", _utc_iso(db.query(func.max(ChatbotAuditRequest.exported_at)).scalar())),
    ]
    for row in rows:
        ws.append(row)
    ws.append(("Failures by stage", "Count"))
    for name, count in db.query(ChatbotAuditStage.stage_name, func.count(ChatbotAuditStage.id)).filter(ChatbotAuditStage.status == "failed").group_by(ChatbotAuditStage.stage_name):
        ws.append((name, count))
    ws.append(("Failures by category", "Count"))
    for name, count in db.query(ChatbotAuditStage.failure_category, func.count(ChatbotAuditStage.id)).filter(ChatbotAuditStage.failure_category.isnot(None)).group_by(ChatbotAuditStage.failure_category):
        ws.append((name, count))
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243B53")


def _export_pending_batch_once():
    """Export a bounded batch; DB rows remain pending if file operations fail."""
    db = SessionLocal()
    lock = None
    batch = []
    export_started = None
    try:
        batch = db.query(ChatbotAuditRequest).filter(ChatbotAuditRequest.exported_at.is_(None)).order_by(ChatbotAuditRequest.created_at).limit(EXPORT_BATCH_SIZE).all()
        if not batch:
            return 0
        export_started = utcnow()
        lock = _lock_file()
        wb = _workbook()
        requests_ws, stages_ws = wb["Request Logs"], wb["Stage Logs"]
        existing = {str(row[0]) for row in requests_ws.iter_rows(min_row=2, max_col=1, values_only=True) if row[0]}
        ids = [row.request_id for row in batch]
        known_stage_keys = {(str(row[0]), str(row[1])) for row in stages_ws.iter_rows(min_row=2, max_col=2, values_only=True) if row[0] and row[1]}
        export_timestamp = utcnow()
        for request in batch:
            export_stage = db.query(ChatbotAuditStage).filter(
                ChatbotAuditStage.request_id == request.request_id,
                ChatbotAuditStage.stage_name == "excel_export",
            ).first()
            if export_stage:
                export_stage.status = "success"
                export_stage.started_at = export_started
                export_stage.completed_at = export_timestamp
                export_stage.duration_ms = (export_timestamp - export_started).total_seconds() * 1000
                export_stage.decision = "exported"
                export_stage.retry_count = request.export_failures or 0
                export_stage.failure_category = None
                export_stage.failure_reason = None
                export_stage.exception_type = None
        stages_by_request = {}
        for stage in db.query(ChatbotAuditStage).filter(ChatbotAuditStage.request_id.in_(ids)).all():
            stages_by_request.setdefault(stage.request_id, []).append(stage)
        for request in batch:
            if request.request_id not in existing:
                values = [request.request_id, _utc_iso(request.created_at), request.session_id, request.user_id, request.query_text,
                          request.answer_text, request.status, request.failure_stage, request.failure_category,
                          request.failure_reason, request.http_status, request.latency_ms, request.source_count,
                          ", ".join(request.source_names or []), request.model_name, request.fallback,
                          request.fallback_reason, _utc_iso(export_timestamp), (request.export_attempts or 0) + 1,
                          request.export_failures or 0, None]
                requests_ws.append([_safe_cell(v) for v in values])
            for stage in stages_by_request.get(request.request_id, []):
                key = (request.request_id, stage.stage_name)
                if key in known_stage_keys:
                    continue
                values = [request.request_id, stage.stage_name, stage.status, _utc_iso(stage.started_at), _utc_iso(stage.completed_at),
                          stage.duration_ms, stage.decision, stage.failure_category, stage.failure_reason,
                          stage.exception_type, stage.retry_count, stage.skip_reason]
                stages_ws.append([_safe_cell(v) for v in values])
                known_stage_keys.add(key)
            existing.add(request.request_id)
        _refresh_summary(wb, db)
        _format_sheet(requests_ws)
        _format_sheet(stages_ws)
        _format_sheet(wb["Summary"])
        temp_path = WORKBOOK_PATH.with_suffix(WORKBOOK_PATH.suffix + ".tmp")
        wb.save(temp_path)
        os.replace(temp_path, WORKBOOK_PATH)
        if os.name != "nt":
            try:
                os.chmod(WORKBOOK_PATH, 0o600)
            except OSError:
                logger.warning("Could not restrict chatbot audit workbook permissions.")
        for row in batch:
            row.exported_at = export_timestamp
            row.export_attempts = (row.export_attempts or 0) + 1
            row.export_error = None
        db.commit()
        return len(batch)
    except Exception as exc:
        db.rollback()
        logger.error("Chatbot audit Excel export failed (%s)", type(exc).__name__)
        for row in batch:
            try:
                row.export_attempts = (row.export_attempts or 0) + 1
                row.export_failures = (row.export_failures or 0) + 1
                row.export_error = type(exc).__name__[:80]
                export_stage = db.query(ChatbotAuditStage).filter(
                    ChatbotAuditStage.request_id == row.request_id,
                    ChatbotAuditStage.stage_name == "excel_export",
                ).first()
                if export_stage:
                    export_stage.status = "failed"
                    export_stage.started_at = export_started or utcnow()
                    export_stage.completed_at = utcnow()
                    export_stage.duration_ms = max(0, (export_stage.completed_at - export_stage.started_at).total_seconds() * 1000)
                    export_stage.failure_category = "excel_export_error"
                    export_stage.failure_reason = "Excel export failed; the database record remains queued for retry."
                    export_stage.exception_type = type(exc).__name__
                    export_stage.retry_count = row.export_attempts
            except Exception:
                pass
        try:
            db.commit()
        except Exception:
            db.rollback()
        return 0
    finally:
        if lock:
            _unlock_file(lock)
        db.close()


def export_pending_batch():
    """Serialize in-process writers as well as the cross-process file lock."""
    with _PROCESS_EXPORT_LOCK:
        return _export_pending_batch_once()


async def export_worker(stop_event):
    """Lifecycle-managed bounded polling worker; no in-memory event queue."""
    while not stop_event.is_set():
        try:
            await asyncio.to_thread(export_pending_batch)
        except Exception as exc:
            logger.error("Chatbot audit export worker failed (%s)", type(exc).__name__)
        try:
            await asyncio.wait_for(stop_event.wait(), timeout=EXPORT_POLL_SECONDS)
        except asyncio.TimeoutError:
            pass
